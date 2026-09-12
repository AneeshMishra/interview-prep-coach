"""
DOCX -> raw text/table extraction.
Keep this dumb on purpose: structuring/interpretation happens in structurer.py.
"""
import hashlib
import re
from dataclasses import dataclass, field

import docx


@dataclass
class RawSection:
    heading: str | None
    paragraphs: list[str] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)


def content_hash(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


# Real interview-notes documents are rarely formatted with Word's built-in
# "Heading N" paragraph styles — in practice people mark a new company/round
# with a short plain-text label instead, e.g. "Nagarro:", "Round 2:",
# "Interview of EPAM on 02/09/2025 1st Round". Relying on Heading styles
# alone means the whole document collapses into one or two giant sections,
# which (a) makes for a huge, fragile single LLM structuring call per
# section, and (b) means one malformed LLM response silently drops every
# question in the document instead of just one company's worth.
#
# This heuristic is intentionally permissive, not precise — it will
# occasionally misfire on a question phrased as an imperative ending in a
# colon (e.g. "Print all permutations of a string:"), splitting it into its
# own section. That's a much smaller loss (one under-sized section instead
# of the right one) than the alternative of never detecting real headings at
# all, and it's still passed through the same LLM structuring step, which is
# what actually has to make sense of the content either way.
_HEADING_MAX_LEN = 90
_COLON_HEADING_RE = re.compile(r"^.{1,%d}:$" % _HEADING_MAX_LEN)
_INTERVIEW_MARKER_RE = re.compile(r"^interview\b", re.IGNORECASE)


def _looks_like_heading(paragraph, text: str) -> bool:
    style_name = paragraph.style.name or ""
    if style_name.lower().startswith("heading"):
        return True
    if not text or len(text) > _HEADING_MAX_LEN or text.endswith("?"):
        return False
    return bool(_COLON_HEADING_RE.match(text) or _INTERVIEW_MARKER_RE.match(text))


def parse_docx(file_path: str) -> list[RawSection]:
    """
    Walk the document splitting on heading paragraphs.
    Anything before the first heading is grouped under heading=None.
    """
    document = docx.Document(file_path)
    sections: list[RawSection] = [RawSection(heading=None)]

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if _looks_like_heading(para, text):
            sections.append(RawSection(heading=text))
        else:
            sections[-1].paragraphs.append(text)

    # Tables aren't anchored to paragraph position in python-docx; attach to
    # the last section as a reasonable default for v1.
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        sections[-1].tables.append(rows)

    return [s for s in sections if s.paragraphs or s.tables]
