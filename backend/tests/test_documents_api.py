import json
import tempfile
from pathlib import Path

import docx
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.config import Settings, get_settings
from app.db.base import Base, get_db
from app.ingestion.google_docs_import import GoogleDocNotAccessible, GoogleDocTooLarge
from app.main import app


class FakeLLM:
    def complete(self, system_prompt: str, user_prompt: str) -> str:
        return json.dumps(
            [
                {
                    "company": "Amazon",
                    "role": "Backend Engineer",
                    "round_type": "system_design",
                    "question": "Design a URL shortener.",
                    "answer_notes": None,
                    "difficulty": "medium",
                    "tags": ["scaling"],
                    "source_type": "user_reported",
                    "source_section": None,
                    "extraction_confidence": 0.9,
                }
            ]
        )


class FakeVectorStore:
    def __init__(self):
        self.upserted = []

    def upsert_question(self, question_id, text, metadata):
        self.upserted.append((question_id, text, metadata))

    def search(self, query, company=None, role=None, round_type=None, limit=20):
        return []


def make_docx_bytes(tmp_path):
    document = docx.Document()
    document.add_heading("Amazon - Backend Engineer", level=1)
    document.add_paragraph("Round 2: Design a URL shortener.")
    path = tmp_path / "sample.docx"
    document.save(path)
    return path.read_bytes()


def make_client(tmp_path, monkeypatch):
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(bind=engine)
    TestingSession = sessionmaker(bind=engine)

    def override_get_db():
        db = TestingSession()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db

    fake_vector_store = FakeVectorStore()
    monkeypatch.setattr("app.ingestion.pipeline.get_llm_provider", lambda settings: FakeLLM())
    monkeypatch.setattr("app.ingestion.pipeline.get_vector_store", lambda: fake_vector_store)

    client = TestClient(app)
    return client, fake_vector_store


def test_upload_ingests_and_questions_are_searchable(tmp_path, monkeypatch):
    client, fake_vector_store = make_client(tmp_path, monkeypatch)
    try:
        file_bytes = make_docx_bytes(tmp_path)

        response = client.post(
            "/api/v1/documents/upload",
            files={"file": ("sample.docx", file_bytes, "application/octet-stream")},
        )
        assert response.status_code == 200
        body = response.json()
        document_id = body["document_id"]

        docs = client.get("/api/v1/documents").json()
        assert docs[0]["id"] == document_id
        assert docs[0]["status"] == "done"

        questions = client.get("/api/v1/questions", params={"company": "Amazon"}).json()
        assert len(questions) == 1
        assert questions[0]["question"] == "Design a URL shortener."
        assert questions[0]["tags"] == ["scaling"]

        assert len(fake_vector_store.upserted) == 1
    finally:
        app.dependency_overrides.clear()


def test_duplicate_upload_is_detected_by_content_hash(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)
    try:
        file_bytes = make_docx_bytes(tmp_path)

        first = client.post(
            "/api/v1/documents/upload",
            files={"file": ("sample.docx", file_bytes, "application/octet-stream")},
        ).json()
        second = client.post(
            "/api/v1/documents/upload",
            files={"file": ("sample.docx", file_bytes, "application/octet-stream")},
        ).json()

        assert second["status"] == "duplicate"
        assert second["document_id"] == first["document_id"]
    finally:
        app.dependency_overrides.clear()


def test_rejects_non_docx_upload(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)
    try:
        response = client.post(
            "/api/v1/documents/upload",
            files={"file": ("sample.txt", b"not a docx", "text/plain")},
        )
        assert response.status_code == 400
    finally:
        app.dependency_overrides.clear()


def test_rejects_upload_over_size_limit(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)
    app.dependency_overrides[get_settings] = lambda: Settings(max_upload_size_bytes=10)
    try:
        file_bytes = make_docx_bytes(tmp_path)
        assert len(file_bytes) > 10

        response = client.post(
            "/api/v1/documents/upload",
            files={"file": ("sample.docx", file_bytes, "application/octet-stream")},
        )
        assert response.status_code == 413

        assert client.get("/api/v1/documents").json() == []
    finally:
        app.dependency_overrides.clear()


def test_upload_sanitizes_path_traversal_filename(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)
    try:
        file_bytes = make_docx_bytes(tmp_path)

        response = client.post(
            "/api/v1/documents/upload",
            files={"file": ("../../etc/passwd.docx", file_bytes, "application/octet-stream")},
        )
        assert response.status_code == 200

        docs = client.get("/api/v1/documents").json()
        assert docs[0]["filename"] == "passwd.docx"
    finally:
        app.dependency_overrides.clear()


def _stub_google_doc_download(file_bytes: bytes, filename: str | None):
    """download_google_doc_as_docx writes to (and the caller deletes) a temp
    file each call, so the stub must produce a fresh one every time too —
    it can't just hand back one path that gets consumed on the first call."""

    async def _download(doc_id, max_bytes, client=None):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(file_bytes)
            tmp_path = Path(tmp.name)
        return tmp_path, filename

    return _download


GOOGLE_DOC_URL = "https://docs.google.com/document/d/1klKw9KYkzLO5L9z0PhnMGpkI2MbMo5wDq9FuzBBNb-E/edit"


def test_import_google_doc_ingests_successfully(tmp_path, monkeypatch):
    client, fake_vector_store = make_client(tmp_path, monkeypatch)
    file_bytes = make_docx_bytes(tmp_path)
    monkeypatch.setattr(
        "app.api.routers.documents.download_google_doc_as_docx",
        _stub_google_doc_download(file_bytes, "Amazon Interview.docx"),
    )
    try:
        response = client.post("/api/v1/documents/import/google-doc", json={"url": GOOGLE_DOC_URL})
        assert response.status_code == 200
        body = response.json()

        docs = client.get("/api/v1/documents").json()
        assert docs[0]["id"] == body["document_id"]
        assert docs[0]["filename"] == "Amazon Interview.docx"
        assert docs[0]["status"] == "done"

        questions = client.get("/api/v1/questions", params={"company": "Amazon"}).json()
        assert len(questions) == 1
        assert len(fake_vector_store.upserted) == 1
    finally:
        app.dependency_overrides.clear()


def test_import_google_doc_falls_back_to_id_based_filename(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)
    file_bytes = make_docx_bytes(tmp_path)
    monkeypatch.setattr(
        "app.api.routers.documents.download_google_doc_as_docx",
        _stub_google_doc_download(file_bytes, None),
    )
    try:
        response = client.post("/api/v1/documents/import/google-doc", json={"url": GOOGLE_DOC_URL})
        assert response.status_code == 200

        docs = client.get("/api/v1/documents").json()
        assert docs[0]["filename"] == "google-doc-1klKw9KYkzLO5L9z0PhnMGpkI2MbMo5wDq9FuzBBNb-E.docx"
    finally:
        app.dependency_overrides.clear()


PUBLISH_TO_WEB_URL = (
    "https://docs.google.com/document/d/e/2PACX-1vREH7wBSxdAMEWhZpuXzzoWWRVFGnawMQ"
    "uSo4JTfPolgT7oWMwq6epoL96_SgtS0_Bw8sieqeQNLYUW/pub"
)


def test_import_google_doc_accepts_publish_to_web_url(tmp_path, monkeypatch):
    """A "Publish to the web" link has a distinct .../d/e/<token>/pub shape;
    the fallback filename must not contain the "/" from that "e/" segment."""
    client, _ = make_client(tmp_path, monkeypatch)
    file_bytes = make_docx_bytes(tmp_path)
    monkeypatch.setattr(
        "app.api.routers.documents.download_google_doc_as_docx",
        _stub_google_doc_download(file_bytes, None),
    )
    try:
        response = client.post("/api/v1/documents/import/google-doc", json={"url": PUBLISH_TO_WEB_URL})
        assert response.status_code == 200

        docs = client.get("/api/v1/documents").json()
        assert "/" not in docs[0]["filename"]
        assert docs[0]["filename"].startswith("google-doc-e-2PACX-")
        assert docs[0]["filename"].endswith(".docx")
    finally:
        app.dependency_overrides.clear()


def test_import_google_doc_detects_duplicate(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)
    file_bytes = make_docx_bytes(tmp_path)
    monkeypatch.setattr(
        "app.api.routers.documents.download_google_doc_as_docx",
        _stub_google_doc_download(file_bytes, "Amazon Interview.docx"),
    )
    try:
        first = client.post("/api/v1/documents/import/google-doc", json={"url": GOOGLE_DOC_URL}).json()
        second = client.post("/api/v1/documents/import/google-doc", json={"url": GOOGLE_DOC_URL}).json()

        assert second["status"] == "duplicate"
        assert second["document_id"] == first["document_id"]
    finally:
        app.dependency_overrides.clear()


def test_import_google_doc_rejects_non_google_docs_url(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)
    try:
        response = client.post(
            "/api/v1/documents/import/google-doc", json={"url": "https://example.com/not-a-doc"}
        )
        assert response.status_code == 400
    finally:
        app.dependency_overrides.clear()


def test_import_google_doc_returns_422_when_not_publicly_accessible(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)

    async def _raise_not_accessible(doc_id, max_bytes, client=None):
        raise GoogleDocNotAccessible("Google Docs did not return a .docx file for this link.")

    monkeypatch.setattr("app.api.routers.documents.download_google_doc_as_docx", _raise_not_accessible)
    try:
        response = client.post("/api/v1/documents/import/google-doc", json={"url": GOOGLE_DOC_URL})
        assert response.status_code == 422
        assert "Anyone with the link" in response.json()["detail"]
    finally:
        app.dependency_overrides.clear()


def test_import_google_doc_returns_413_when_export_too_large(tmp_path, monkeypatch):
    client, _ = make_client(tmp_path, monkeypatch)

    async def _raise_too_large(doc_id, max_bytes, client=None):
        raise GoogleDocTooLarge(f"Export exceeds maximum upload size of {max_bytes} bytes.")

    monkeypatch.setattr("app.api.routers.documents.download_google_doc_as_docx", _raise_too_large)
    try:
        response = client.post("/api/v1/documents/import/google-doc", json={"url": GOOGLE_DOC_URL})
        assert response.status_code == 413
    finally:
        app.dependency_overrides.clear()
