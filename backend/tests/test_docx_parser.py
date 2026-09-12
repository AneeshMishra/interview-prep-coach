import io

import docx
import pytest

from app.ingestion.docx_parser import content_hash, parse_docx


def make_docx(tmp_path, paragraphs_with_headings):
    document = docx.Document()
    for text, is_heading in paragraphs_with_headings:
        if is_heading:
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    path = tmp_path / "sample.docx"
    document.save(path)
    return path


def test_parse_docx_splits_on_headings(tmp_path):
    path = make_docx(
        tmp_path,
        [
            ("Amazon - Backend Engineer", True),
            ("Round 1: Tell me about yourself.", False),
            ("Round 2: System Design", True),
            ("Design a URL shortener.", False),
        ],
    )
    sections = parse_docx(str(path))
    headings = [s.heading for s in sections]
    assert "Amazon - Backend Engineer" in headings
    assert "Round 2: System Design" in headings


def test_parse_docx_treats_short_colon_terminated_lines_as_headings(tmp_path):
    # Real interview notes are rarely formatted with Word's Heading style —
    # people mark a new company with a plain "Nagarro:"-style label instead.
    path = make_docx(
        tmp_path,
        [
            ("Nagarro:", False),
            ("What is virtual thread", False),
            ("Serigornic:", False),
            ("Explain Kafka design principles", False),
        ],
    )
    sections = parse_docx(str(path))
    headings = [s.heading for s in sections]
    assert "Nagarro:" in headings
    assert "Serigornic:" in headings

    nagarro = next(s for s in sections if s.heading == "Nagarro:")
    assert nagarro.paragraphs == ["What is virtual thread"]


def test_parse_docx_treats_interview_prefixed_lines_as_headings(tmp_path):
    path = make_docx(
        tmp_path,
        [
            ("Interview of EPAM on 02/09/2025 1st Round", False),
            ("What is a marker interface?", False),
        ],
    )
    sections = parse_docx(str(path))
    headings = [s.heading for s in sections]
    assert "Interview of EPAM on 02/09/2025 1st Round" in headings


def test_parse_docx_does_not_treat_questions_as_headings(tmp_path):
    path = make_docx(
        tmp_path,
        [
            ("Nagarro:", False),
            ("What is the interview process for this role?", False),
        ],
    )
    sections = parse_docx(str(path))
    nagarro = next(s for s in sections if s.heading == "Nagarro:")
    # Ends with "?" and isn't short-and-colon-terminated, so it stays
    # grouped under the preceding heading rather than starting a new one.
    assert nagarro.paragraphs == ["What is the interview process for this role?"]


def test_parse_docx_does_not_treat_long_colon_terminated_lines_as_headings(tmp_path):
    long_line = (
        "This is a long sentence that happens to end with a colon so it should "
        "not be mistaken for a section heading even though it ends with a colon:"
    )
    path = make_docx(tmp_path, [("Nagarro:", False), (long_line, False)])
    sections = parse_docx(str(path))
    nagarro = next(s for s in sections if s.heading == "Nagarro:")
    assert nagarro.paragraphs == [long_line]


def test_content_hash_is_deterministic():
    data = b"same bytes"
    assert content_hash(data) == content_hash(data)


def test_content_hash_differs_for_different_content():
    assert content_hash(b"a") != content_hash(b"b")
