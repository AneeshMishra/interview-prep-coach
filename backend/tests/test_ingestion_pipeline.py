import json

import docx
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db.base import Base
from app.db.models import Document, InterviewQuestion, QuestionTag
from app.ingestion.pipeline import run_ingestion


class FakeLLM:
    """Returns one canned structured record per section, ignoring content."""

    def complete(self, system_prompt: str, user_prompt: str) -> str:
        return json.dumps(
            [
                {
                    "company": "Amazon",
                    "role": "Backend Engineer",
                    "round_type": "system_design",
                    "question": "Design a URL shortener.",
                    "answer_notes": "Discuss hashing and sharding.",
                    "difficulty": "medium",
                    "tags": ["Scaling", "  Hashing "],
                    "source_type": "user_reported",
                    "source_section": None,
                    "extraction_confidence": 0.9,
                }
            ]
        )


class FailingLLM:
    def complete(self, system_prompt: str, user_prompt: str) -> str:
        raise RuntimeError("LLM unreachable")


class FakeVectorStore:
    def __init__(self):
        self.upserted = []

    def upsert_question(self, question_id, text, metadata):
        self.upserted.append((question_id, text, metadata))


@pytest.fixture
def db_session(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    yield session
    session.close()


def make_docx(tmp_path):
    document = docx.Document()
    document.add_heading("Amazon - Backend Engineer", level=1)
    document.add_paragraph("Round 2: Design a URL shortener.")
    path = tmp_path / "sample.docx"
    document.save(path)
    return path


def test_run_ingestion_persists_questions_tags_and_embeddings(db_session, tmp_path):
    doc_path = make_docx(tmp_path)
    document = Document(filename="sample.docx", content_hash="hash1", status="pending")
    db_session.add(document)
    db_session.commit()

    vector_store = FakeVectorStore()
    run_ingestion(document.id, str(doc_path), db_session, llm=FakeLLM(), vector_store=vector_store)

    db_session.refresh(document)
    assert document.status == "done"
    assert document.error_message is None

    questions = db_session.query(InterviewQuestion).filter_by(document_id=document.id).all()
    assert len(questions) == 1
    question = questions[0]
    assert question.company == "Amazon"
    assert question.round_type == "system_design"

    tags = db_session.query(QuestionTag).filter_by(question_id=question.id).all()
    assert {t.tag for t in tags} == {"scaling", "hashing"}

    assert len(vector_store.upserted) == 1
    upserted_id, upserted_text, metadata = vector_store.upserted[0]
    assert upserted_id == question.id
    assert upserted_text == "Design a URL shortener."
    assert metadata["company"] == "Amazon"

    # temp docx should be cleaned up after ingestion
    assert not doc_path.exists()


def test_run_ingestion_marks_document_failed_on_llm_error(db_session, tmp_path):
    doc_path = make_docx(tmp_path)
    document = Document(filename="sample.docx", content_hash="hash2", status="pending")
    db_session.add(document)
    db_session.commit()

    run_ingestion(document.id, str(doc_path), db_session, llm=FailingLLM(), vector_store=FakeVectorStore())

    db_session.refresh(document)
    assert document.status == "failed"
    assert "LLM unreachable" in document.error_message

    assert db_session.query(InterviewQuestion).filter_by(document_id=document.id).count() == 0
