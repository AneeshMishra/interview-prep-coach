"""
Endpoint-level proof that rate limiting (app/rate_limit.py) is actually
wired into the right routes with the right keying — auth by IP, uploads
and LLM-calling endpoints by signed-in user. Each test overrides only the
one limiter getter it cares about with a deliberately tiny limit, via the
same app.dependency_overrides mechanism used throughout this test suite
(get_db, get_settings, ...) — the autouse fixture in conftest.py otherwise
keeps every other test's limiter effectively unlimited.
"""
import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.base import Base, get_db
from app.main import app
from app.rate_limit import RateLimiter, get_auth_rate_limiter, get_llm_rate_limiter, get_upload_rate_limiter
from tests.auth_helpers import authenticate, create_user


@pytest.fixture
def db_client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    TestingSession = sessionmaker(bind=engine)

    def override_get_db():
        db = TestingSession()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    yield TestClient(app), TestingSession
    app.dependency_overrides.clear()


def make_docx_bytes() -> bytes:
    import io

    document = docx.Document()
    document.add_heading("Amazon - Backend Engineer", level=1)
    document.add_paragraph("Round 2: Design a URL shortener.")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_login_is_rate_limited_by_ip(db_client):
    client, _ = db_client
    auth_limiter = RateLimiter(max_requests=2, window_seconds=60)
    app.dependency_overrides[get_auth_rate_limiter] = lambda: auth_limiter

    first = client.get("/api/v1/auth/google/login", follow_redirects=False)
    second = client.get("/api/v1/auth/google/login", follow_redirects=False)
    third = client.get("/api/v1/auth/google/login", follow_redirects=False)

    assert first.status_code != 429
    assert second.status_code != 429
    assert third.status_code == 429
    assert "Retry-After" in third.headers


def test_auth_rate_limit_bucket_is_shared_across_login_and_refresh_for_one_ip(db_client):
    """Login, callback and refresh are all part of the same auth attack
    surface from one source IP — they intentionally share one bucket per
    IP rather than each getting their own independent quota."""
    client, _ = db_client
    auth_limiter = RateLimiter(max_requests=2, window_seconds=60)
    app.dependency_overrides[get_auth_rate_limiter] = lambda: auth_limiter

    client.get("/api/v1/auth/google/login", follow_redirects=False)  # 1/2
    client.post("/api/v1/auth/refresh")  # 2/2 (401 — no cookie — but still counts)
    third = client.get("/api/v1/auth/google/login", follow_redirects=False)

    assert third.status_code == 429


def test_upload_rate_limit_is_scoped_per_user(db_client):
    client, TestingSession = db_client
    upload_limiter = RateLimiter(max_requests=1, window_seconds=60)
    app.dependency_overrides[get_upload_rate_limiter] = lambda: upload_limiter

    alice_id = create_user(TestingSession(), email="alice@example.com").id
    bob_id = create_user(TestingSession(), email="bob@example.com").id
    file_bytes = make_docx_bytes()

    authenticate(client, alice_id)
    first = client.post(
        "/api/v1/documents/upload", files={"file": ("a.docx", file_bytes, "application/octet-stream")}
    )
    assert first.status_code != 429
    second = client.post(
        "/api/v1/documents/upload", files={"file": ("b.docx", file_bytes, "application/octet-stream")}
    )
    assert second.status_code == 429

    # Bob has never uploaded — his own quota is untouched by Alice's usage.
    authenticate(client, bob_id)
    third = client.post(
        "/api/v1/documents/upload", files={"file": ("c.docx", file_bytes, "application/octet-stream")}
    )
    assert third.status_code != 429


def test_upload_rate_limit_also_applies_to_google_doc_import(db_client, monkeypatch):
    client, TestingSession = db_client
    upload_limiter = RateLimiter(max_requests=1, window_seconds=60)
    app.dependency_overrides[get_upload_rate_limiter] = lambda: upload_limiter

    async def _fail_download(doc_id, max_bytes, client=None):
        from app.ingestion.google_docs_import import GoogleDocNotAccessible

        raise GoogleDocNotAccessible("not accessible")

    monkeypatch.setattr("app.api.routers.documents.download_google_doc_as_docx", _fail_download)

    user_id = create_user(TestingSession()).id
    authenticate(client, user_id)
    url = "https://docs.google.com/document/d/abc123/edit"

    first = client.post("/api/v1/documents/import/google-doc", json={"url": url})
    assert first.status_code != 429  # 422 (not accessible), but not rate-limited
    second = client.post("/api/v1/documents/import/google-doc", json={"url": url})
    assert second.status_code == 429


def test_chat_message_rate_limit_is_scoped_per_user(db_client):
    client, TestingSession = db_client
    llm_limiter = RateLimiter(max_requests=1, window_seconds=60)
    app.dependency_overrides[get_llm_rate_limiter] = lambda: llm_limiter

    alice_id = create_user(TestingSession(), email="alice@example.com").id
    bob_id = create_user(TestingSession(), email="bob@example.com").id

    authenticate(client, alice_id)
    alice_session = client.post("/api/v1/chat/sessions").json()
    first = client.post(f"/api/v1/chat/sessions/{alice_session['id']}/messages", json={"message": "hi"})
    assert first.status_code != 429
    second = client.post(f"/api/v1/chat/sessions/{alice_session['id']}/messages", json={"message": "hi again"})
    assert second.status_code == 429

    authenticate(client, bob_id)
    bob_session = client.post("/api/v1/chat/sessions").json()
    third = client.post(f"/api/v1/chat/sessions/{bob_session['id']}/messages", json={"message": "hi"})
    assert third.status_code != 429


def test_llm_rate_limit_bucket_is_shared_between_chat_and_interviews_for_one_user(db_client):
    """Chat and mock interviews both call an LLM on the app's behalf — they
    intentionally share one per-user LLM-usage bucket rather than each
    getting an independent quota, so the real cost driver (LLM calls) is
    capped regardless of which feature triggers them."""
    client, TestingSession = db_client
    llm_limiter = RateLimiter(max_requests=1, window_seconds=60)
    app.dependency_overrides[get_llm_rate_limiter] = lambda: llm_limiter

    user_id = create_user(TestingSession()).id
    authenticate(client, user_id)

    session = client.post("/api/v1/chat/sessions").json()
    client.post(f"/api/v1/chat/sessions/{session['id']}/messages", json={"message": "hi"})  # consumes the 1 slot

    response = client.post("/api/v1/interviews", json={"company": "Amazon"})
    assert response.status_code == 429
